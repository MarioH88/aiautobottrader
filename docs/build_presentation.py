"""Generate a clean PowerPoint project completion presentation."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

try:
    from PIL import Image
except ImportError:
    Image = None  # type: ignore

DOCS = Path(__file__).resolve().parent
ASSETS = DOCS / "report_assets"
OUTPUT = DOCS / "AI_AutoBot_Trader_Progress_Report_v2.pptx"

ACCENT = RGBColor(79, 70, 229)
DARK = RGBColor(15, 23, 42)
WHITE = RGBColor(255, 255, 255)
CODE_BG = RGBColor(15, 23, 42)


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _add_header(slide, title: str, subtitle: str = "") -> None:
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.55))
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = ACCENT

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.78), Inches(9), Inches(0.35))
        stf = sub_box.text_frame
        stf.text = subtitle
        sp = stf.paragraphs[0]
        sp.font.size = Pt(14)
        sp.font.color.rgb = RGBColor(100, 116, 139)


def _add_bullets(slide, top: float, items: list[str], size: int = 16) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(8.9), Inches(6.2 - top))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(6)


def _add_code_block(slide, top: float, code: str, height: float = 2.8) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(8.9), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = CODE_BG
    box.line.color.rgb = RGBColor(51, 65, 85)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.15)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    p.text = code
    p.font.name = "Consolas"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(226, 232, 240)


def _add_table(slide, top: float, rows: list[tuple[str, str]], font_size: int = 12) -> None:
    row_count = len(rows)
    height = min(Inches(0.32 * row_count), Inches(5.6))
    table_shape = slide.shapes.add_table(row_count, 2, Inches(0.5), Inches(top), Inches(9), height)
    table = table_shape.table
    table.columns[0].width = Inches(6.4)
    table.columns[1].width = Inches(2.1)
    for r, (col1, col2) in enumerate(rows):
        table.cell(r, 0).text = col1
        table.cell(r, 1).text = col2
        for c in range(2):
            cell = table.cell(r, c)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(font_size + (2 if r == 0 else 0))
                paragraph.font.bold = r == 0 or (r == row_count - 1 and col1 == "OVERALL PROJECT")
                if r == 0:
                    paragraph.font.color.rgb = ACCENT
                elif r == row_count - 1:
                    paragraph.font.color.rgb = ACCENT


def _add_picture_fit(slide, path: Path, top: float = 1.15, max_width: float = 8.8, max_height: float = 5.9) -> None:
    if not path.exists():
        return
    if Image is None:
        slide.shapes.add_picture(str(path), Inches(0.6), Inches(top), width=Inches(max_width))
        return
    with Image.open(path) as img:
        px_w, px_h = img.size
    aspect = px_h / px_w
    width = max_width
    height = width * aspect
    if height > max_height:
        height = max_height
        width = height / aspect
    left = (10 - width) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # --- Slide 1: Title ---
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "AI AutoBot Trader"
    slide.placeholders[1].text = (
        "Project Completion Report (100%)\n"
        "Author: Mario Hayden\n"
        "Date: July 27, 2026\n"
        "Overall Completion: 100%"
    )
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(40)
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = ACCENT
    for p in slide.placeholders[1].text_frame.paragraphs:
        p.font.size = Pt(18)

    # --- Slide 2: Description ---
    slide = _blank_slide(prs)
    _add_header(slide, "Project Description", "Automated algorithmic trading with ML and sentiment")
    _add_bullets(
        slide,
        1.15,
        [
            "Processes live and historical OHLCV (stocks + crypto) via Alpaca",
            "Ingests financial news via Alpaca NewsClient and scores with VADER",
            "Loads macro sentiment from a Kaggle-style CSV export",
            "Uses Random Forest / XGBoost plus RSI, MACD, and ATR for signals",
            "Executes trades with dynamic / fractional position sizing",
            "Manages risk with broker-backed ATR trailing stop-losses",
            "Includes Streamlit dashboard, scheduler, Docker, and deploy scripts",
        ],
        size=15,
    )

    # --- Slide 3: Architecture ---
    slide = _blank_slide(prs)
    _add_header(slide, "System Architecture", "Async event-driven modular design")
    _add_code_block(
        slide,
        1.15,
        "main.py\n"
        "  -> bot/trading_bot.py\n"
        "       data_pipeline.py       (stock/crypto OHLCV)\n"
        "       news_sentiment.py      (Alpaca News + VADER)\n"
        "       macro_sentiment.py     (Kaggle CSV)\n"
        "       strategy_engine.py     (ML signals)\n"
        "       allocation_engine.py   (fractional sizing)\n"
        "       execution_client.py    (orders + broker stops)\n"
        "       stream_client.py       (stock/crypto WebSocket)\n"
        "       scheduler.py           (trading windows)\n"
        "       connection_manager.py  (retry / recovery)\n"
        "  dashboard_app.py           (Streamlit UI)",
        height=4.8,
    )

    # --- Slides 4-5: Feature tables (split to avoid overflow) ---
    rows_a = [
        ("Feature / Task", "Done"),
        ("Modular async bot architecture", "100%"),
        ("Alpaca REST OHLCV pipeline", "100%"),
        ("WebSocket live streaming", "100%"),
        ("Technical indicators (RSI/MACD/ATR)", "100%"),
        ("News sentiment (VADER + crypto aliases)", "100%"),
        ("Macro sentiment (Kaggle CSV)", "100%"),
        ("ML training and inference", "100%"),
        ("Order execution and allocation", "100%"),
    ]
    slide = _blank_slide(prs)
    _add_header(slide, "Feature Completion (Part 1)", "Core trading pipeline")
    _add_table(slide, 1.15, rows_a)

    rows_b = [
        ("Feature / Task", "Done"),
        ("ATR trailing stops (broker + software)", "100%"),
        ("Auto-start trading (AUTO_START)", "100%"),
        ("Trading window scheduler", "100%"),
        ("Streamlit dashboard UI", "100%"),
        ("Live market prices", "100%"),
        ("Manual buy/sell UI", "100%"),
        ("Backtesting", "100%"),
        ("Integration / unit tests", "100%"),
        ("Crypto OHLCV support", "100%"),
        ("Production deployment", "100%"),
        ("OVERALL PROJECT", "100%"),
    ]
    slide = _blank_slide(prs)
    _add_header(slide, "Feature Completion (Part 2)", "UI, testing, crypto, and deployment")
    _add_table(slide, 1.15, rows_b)

    # --- Slide 6: UI ---
    slide = _blank_slide(prs)
    _add_header(slide, "User Interface Screenshot", "Streamlit dashboard with live prices and manual trade")
    _add_picture_fit(slide, ASSETS / "dashboard_ui_screenshot.png", top=1.05, max_height=5.8)

    # --- Slide 7: Code screenshot ---
    slide = _blank_slide(prs)
    _add_header(slide, "Code Screenshot", "main.py - schedule-aware auto-start")
    _add_picture_fit(slide, ASSETS / "code_screenshot_main.png", top=1.05, max_height=5.8)

    # --- Slide 8: Code listing ---
    slide = _blank_slide(prs)
    _add_header(slide, "Code Listing", "bot/trading_bot.py - execution loop")
    _add_code_block(
        slide,
        1.15,
        "async def execution_tick(self, symbol):\n"
        "    sentiment, macro = self._combined_sentiment(symbol)\n"
        "    frame = await self.data_pipeline.fetch_bars(symbol)\n"
        "    bar_state = self.data_pipeline.latest_bar_state(...)\n"
        "    signal = self.strategy_engine.evaluate(bar_state)\n"
        "    await self.execution_client.process_signal(signal, bar_state)",
        height=3.2,
    )

    # --- Slide 9: How the bot learns ---
    slide = _blank_slide(prs)
    _add_header(slide, "How the Bot Learns to Trade", "Offline training, then live inference")
    _add_bullets(
        slide,
        1.15,
        [
            "1. Training (offline): python train_model.py reads train_data.csv",
            "2. Features: RSI, MACD, and news/macro sentiment_score",
            "3. Target: short-term price direction (up = 1, down = 0)",
            "4. Model: Random Forest (default) or XGBoost → saved as model.pkl",
            "5. Live trading: bot computes the same features and asks model.pkl for buy/sell/hold",
            "6. Allocation sizes the order; ATR trailing stops manage risk",
            "7. Fallback: if the model is missing, RSI/MACD rules still trade",
            "Note: the bot does not auto-retrain after each trade — retrain manually to improve",
        ],
        size=15,
    )

    # --- Slide 10: Tech stack ---
    slide = _blank_slide(prs)
    _add_header(slide, "Technology Stack", "Frameworks and libraries used")
    _add_bullets(
        slide,
        1.15,
        [
            "alpaca-py - market data, news, and trading API",
            "pandas / numpy - feature engineering",
            "scikit-learn / XGBoost - ML direction classifier",
            "pandas-ta - RSI, MACD, ATR (with pandas fallback)",
            "pydantic - validated market and position schemas",
            "vaderSentiment - news headline scoring",
            "Streamlit - web dashboard UI",
            "python-dotenv - secure API credential management",
            "backtrader - historical strategy backtesting",
        ],
        size=15,
    )

    # --- Slide 11: Challenges ---
    slide = _blank_slide(prs)
    _add_header(slide, "Challenges Encountered", "Problems faced and how they were resolved")
    _add_bullets(
        slide,
        1.15,
        [
            "Alpaca WebSocket connection limit when multiple instances ran",
            "  -> Fixed with REST polling and a reconnect circuit breaker",
            "ML models overfit on historical data vs live markets",
            "  -> Added sentiment features and rule-based fallback",
            "Fractional share orders rejected (GTC not allowed)",
            "  -> Switched fractional orders to DAY time-in-force",
            "Windows wrote bot flag file with UTF-8 BOM",
            "  -> Built UTF-8-safe flag file reader/writer",
            "pandas-ta broke on NumPy 2.x",
            "  -> Added pure-pandas indicator fallback",
        ],
        size=14,
    )

    # --- Slide 12: How to run ---
    slide = _blank_slide(prs)
    _add_header(slide, "How to Run", "Commands to verify the project")
    _add_code_block(
        slide,
        1.15,
        "cd aiautobottrader\n"
        "pip install -r requirements.txt\n"
        "python train_model.py\n"
        "python main.py\n"
        "streamlit run dashboard_app.py\n"
        "python test_bot.py\n"
        "python backtest.py\n"
        "docker compose up -d --build",
        height=3.8,
    )
    _add_bullets(
        slide,
        5.15,
        [
            "Status: crypto, scheduler, Docker deploy, and broker ATR stops complete (100%)",
        ],
        size=14,
    )

    # --- Slide 13: Thank you ---
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Thank You"
    slide.placeholders[1].text = "AI AutoBot Trader\nMario Hayden\n100% Complete"
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = ACCENT
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(44)
    for p in slide.placeholders[1].text_frame.paragraphs:
        p.font.size = Pt(20)
        p.alignment = PP_ALIGN.CENTER

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Created: {path}")
